from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app import create_app, db
from models.attendance import Attendance
from models.academic import Department
from models.audit import SecurityEvent
from models.faculty import Faculty
from models.notifications import Notification, NotificationPreference
from models.prediction import PredictionHistory
from models.student import Student
from models.subject import Subject
from models.user import User
from security import hash_password


def _seed(app):
    with app.app_context():
        cse = Department(code='CSE', name='Computer Science')
        it = Department(code='IT', name='Information Technology')
        db.session.add_all([cse, it])
        db.session.flush()
        admin = User(username='report-admin', password_hash=hash_password('password'), role='SUPER_ADMIN')
        management = User(username='report-management', password_hash=hash_password('password'), role='MANAGEMENT')
        cse_user = User(username='report-cse-student', password_hash=hash_password('password'), role='STUDENT', department_id=cse.id)
        it_user = User(username='report-it-student', password_hash=hash_password('password'), role='STUDENT', department_id=it.id)
        cse_faculty_user = User(username='report-cse-faculty', password_hash=hash_password('password'), role='FACULTY', department_id=cse.id)
        it_faculty_user = User(username='report-it-faculty', password_hash=hash_password('password'), role='FACULTY', department_id=it.id)
        cse_hod = User(username='report-cse-hod', password_hash=hash_password('password'), role='HOD', department_id=cse.id)
        cse_department_admin = User(
            username='report-cse-dept-admin', password_hash=hash_password('password'),
            role='DEPARTMENT_ADMIN', department_id=cse.id,
        )
        db.session.add_all([admin, management, cse_user, it_user, cse_faculty_user, it_faculty_user, cse_hod, cse_department_admin])
        db.session.flush()
        cse_student = Student(
            register_number='CSE-10', name='CSE Student', email='cse@example.test', phone='10',
            department='CSE', department_id=cse.id, year='1', section='A', semester='1', user_id=cse_user.id,
        )
        it_student = Student(
            register_number='IT-10', name='IT Student', email='it@example.test', phone='11',
            department='IT', department_id=it.id, year='1', section='A', semester='1', user_id=it_user.id,
        )
        cse_faculty = Faculty(
            employee_id='CSE-10', name='CSE Faculty', email='cf@example.test', department='CSE',
            department_id=cse.id, user_id=cse_faculty_user.id,
        )
        it_faculty = Faculty(
            employee_id='IT-10', name='IT Faculty', email='if@example.test', department='IT',
            department_id=it.id, user_id=it_faculty_user.id,
        )
        db.session.add_all([cse_student, it_student, cse_faculty, it_faculty])
        db.session.flush()
        cse_subject = Subject(
            subject_code='CSE-10', subject_name='Algorithms', department='CSE', semester='1', faculty_id=cse_faculty.id,
        )
        it_subject = Subject(
            subject_code='IT-10', subject_name='Networks', department='IT', semester='1', faculty_id=it_faculty.id,
        )
        db.session.add_all([cse_subject, it_subject])
        db.session.flush()
        db.session.add_all([
            Attendance(student_id=cse_student.id, subject_id=cse_subject.id, date=date(2026, 10, 1), status='PRESENT'),
            Attendance(student_id=cse_student.id, subject_id=cse_subject.id, date=date(2026, 10, 2), status='ABSENT'),
            Attendance(student_id=it_student.id, subject_id=it_subject.id, date=date(2026, 10, 1), status='PRESENT'),
        ])
        db.session.commit()
        return {
            'admin': admin.username, 'management': management.username,
            'cse_student_user': cse_user.username, 'it_student_user': it_user.username,
            'cse_faculty': cse_faculty_user.username, 'it_faculty': it_faculty_user.username,
            'hod': cse_hod.username, 'cse_student_id': cse_student.id,
            'department_admin': cse_department_admin.username,
            'it_student_id': it_student.id, 'cse_user_id': cse_user.id, 'it_user_id': it_user.id,
        }


def _client(app, username):
    client = app.test_client()
    response = client.post('/login', data={'username': username, 'password': 'password'})
    assert response.status_code == 302
    return client


def test_prediction_api_persists_real_attendance_and_safe_history():
    app = create_app(testing=True)
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    ids = _seed(app)
    with app.app_context():
        student_id = db.session.query(Student.id).filter_by(user_id=User.query.filter_by(username=ids['cse_student_user']).one().id).scalar()
    client = _client(app, ids['cse_student_user'])
    response = client.get(f'/api/prediction/student/{student_id}')
    assert response.status_code == 200
    payload = response.get_json()
    assert payload['student'] == 'CSE Student'
    assert payload['current_percentage'] == 50.0
    assert payload['predicted_percentage'] == 50.0
    assert 'confidence' not in payload and 'debug' not in payload
    with app.app_context():
        history = PredictionHistory.query.filter_by(student_id=student_id).one()
        assert history.model_version == 'deterministic-v1'
        assert history.feature_snapshot == {
            'attendance_records': 2, 'present_records': 1, 'recent_window': 2,
            'recent_present_records': 1, 'recent_percentage': 50.0,
        }
        assert history.created_at is not None
    history_response = client.get(f'/api/prediction/student/{student_id}/history')
    assert history_response.status_code == 200
    assert len(history_response.get_json()) == 1
    assert client.get(f"/api/prediction/student/{ids['it_student_id']}/history").status_code == 403
    assert client.get('/student/prediction').status_code == 200
    with app.app_context():
        assert PredictionHistory.query.filter_by(student_id=student_id).count() == 2


def test_notifications_are_owner_scoped_and_preferences_are_user_specific():
    app = create_app(testing=True)
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    ids = _seed(app)
    with app.app_context():
        own = Notification(user_id=ids['cse_user_id'], category='leave', title='Own', message='Own message')
        other = Notification(user_id=ids['it_user_id'], category='leave', title='Other', message='Other message')
        db.session.add_all([own, other])
        db.session.commit()
        own_id, other_id = own.id, other.id
    first_client = _client(app, ids['cse_student_user'])
    second_client = _client(app, ids['it_student_user'])
    listing = first_client.get('/api/notifications').get_json()['notifications']
    assert [item['id'] for item in listing] == [own_id]
    assert first_client.post(f'/api/notifications/{other_id}/read').status_code == 404
    assert first_client.post(f'/api/notifications/{other_id}/dismiss').status_code == 404
    assert first_client.post(f'/api/notifications/{own_id}/read').status_code == 200
    assert first_client.post(f'/api/notifications/{own_id}/dismiss').status_code == 200
    assert first_client.get('/api/notifications').get_json()['notifications'] == []
    assert first_client.get('/api/notification-preferences').get_json() == {
        'leave_updates': True, 'correction_updates': True, 'security_alerts': True,
    }
    saved = first_client.post('/api/notification-preferences', json={'leave_updates': False})
    assert saved.status_code == 200 and saved.get_json()['leave_updates'] is False
    assert second_client.get('/api/notification-preferences').get_json()['leave_updates'] is True
    assert first_client.get('/notifications').status_code == 200
    assert first_client.get('/notification-preferences').status_code == 200
    with app.app_context():
        assert db.session.get(Notification, own_id).read_at is not None
        assert db.session.get(Notification, own_id).dismissed_at is not None
        assert db.session.get(NotificationPreference, ids['cse_user_id']).leave_updates is False


def test_attendance_reports_are_scoped_formatted_and_audited():
    app = create_app(testing=True)
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    ids = _seed(app)

    student = _client(app, ids['cse_student_user'])
    xlsx = student.get('/api/export/attendance.xlsx')
    assert xlsx.status_code == 200
    assert xlsx.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    sheet = load_workbook(BytesIO(xlsx.data), read_only=True).active
    content = list(sheet.values)
    assert len(content) == 4  # title, header, and the student's two attendance rows
    assert all('IT Student' not in str(row) for row in content)
    assert student.get('/api/export/attendance.pdf').data.startswith(b'%PDF')

    faculty = _client(app, ids['cse_faculty'])
    faculty_export = faculty.get('/api/export/attendance.xlsx')
    faculty_sheet = load_workbook(BytesIO(faculty_export.data), read_only=True).active
    faculty_rows = list(faculty_sheet.values)
    assert len(faculty_rows) == 2  # title and header; legacy rows lack a matching assigned class session
    assert all('IT Student' not in str(row) for row in faculty_rows)

    hod = _client(app, ids['hod'])
    assert hod.get('/api/export/attendance.pdf').status_code == 200
    department_admin = _client(app, ids['department_admin'])
    dept_rows = list(load_workbook(
        BytesIO(department_admin.get('/api/export/attendance.xlsx').data), read_only=True,
    ).active.values)
    assert len(dept_rows) == 4
    assert all('IT Student' not in str(row) for row in dept_rows)
    management = _client(app, ids['management'])
    management_export = management.get('/api/export/attendance.xlsx')
    management_rows = list(load_workbook(BytesIO(management_export.data), read_only=True).active.values)
    assert len(management_rows) == 4  # title, header, CSE and IT aggregates
    assert all('CSE Student' not in str(row) and 'IT Student' not in str(row) for row in management_rows)
    assert management.get('/api/export/attendance.pdf').status_code == 200

    administrator = _client(app, ids['admin'])
    admin_export = administrator.get('/api/export/attendance.xlsx')
    admin_rows = list(load_workbook(BytesIO(admin_export.data), read_only=True).active.values)
    assert len(admin_rows) == 5
    assert administrator.get('/api/export/attendance.csv').status_code == 200
    with app.app_context():
        assert SecurityEvent.query.filter(SecurityEvent.action.like('%export_attendance_file%')).count() >= 5
        formats = {
            event.resource_id for event in SecurityEvent.query.filter_by(resource_type='export').all()
        }
        assert {'xlsx', 'pdf'}.issubset(formats)


def test_repeated_failed_signins_notify_super_admins_and_respect_preferences():
    app = create_app(testing=True)
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    with app.app_context():
        administrator = User(username='security-alert-admin', password_hash=hash_password('password'), role='SUPER_ADMIN')
        db.session.add(administrator)
        db.session.commit()
        administrator_id = administrator.id
    client = app.test_client()
    for _ in range(5):
        assert client.post('/login', data={
            'username': 'unknown-alert-account', 'password': 'invalid',
        }).status_code == 200
    with app.app_context():
        alert = Notification.query.filter_by(user_id=administrator_id, category='security').one()
        assert alert.title == 'Repeated failed sign-ins'

    muted_app = create_app(testing=True)
    muted_app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    with muted_app.app_context():
        muted_admin = User(username='muted-security-admin', password_hash=hash_password('password'), role='SUPER_ADMIN')
        db.session.add(muted_admin)
        db.session.flush()
        db.session.add(NotificationPreference(user_id=muted_admin.id, security_alerts=False))
        db.session.commit()
    muted_client = muted_app.test_client()
    for _ in range(5):
        assert muted_client.post('/login', data={
            'username': 'unknown-muted-account', 'password': 'invalid',
        }).status_code == 200
    with muted_app.app_context():
        assert Notification.query.filter_by(category='security').count() == 0