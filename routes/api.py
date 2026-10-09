from datetime import datetime, timedelta
import csv
from io import StringIO

from flask import Blueprint, Response, current_app, jsonify, request, url_for
from flask_login import current_user, login_required

from extensions import db
from attendance_engine import (
    DEFAULT_TARGET_PERCENTAGE, calculate_percentage, required_classes_to_attend, risk_level,
)
from models.academic import ClassSession, Department, Section
from models.attendance import Attendance, AttendanceCorrectionRequest
from models.faculty import Faculty
from models.student import Student
from models.subject import Subject
from models.prediction import PredictionHistory
from permissions import Role, can_access, department_id_for, has_role, role_for
from prediction_history import record_student_prediction
from attendance_imports import (
    ImportValidationError, commit_import_stage, create_import_stage,
    expire_stage, import_preview_data, validate_csv_upload,
)
from models.operational import AttendanceImportStage

api_bp = Blueprint('api', __name__)


def _student_department_scope(student):
    if student.department_id is not None:
        return student.department_id
    value = str(student.department or '').strip()
    if not value:
        return None
    matches = Department.query.filter(db.or_(
        db.func.lower(db.func.trim(Department.code)) == value.casefold(),
        db.func.lower(db.func.trim(Department.name)) == value.casefold(),
    )).all()
    if len(matches) > 1:
        return None
    return matches[0].id if matches else value


def _scope_department_values(department_scope):
    if department_scope is None:
        return set()
    try:
        department = db.session.get(Department, int(department_scope))
    except (TypeError, ValueError):
        department = None
    if department:
        return {department.code.casefold(), department.name.casefold()}
    return {str(department_scope).strip().casefold()}


def _students_for_department(department_scope):
    values = _scope_department_values(department_scope)
    if not values:
        return Student.query.filter(Student.id == -1)
    conditions = [Student.department_id == department_scope] if isinstance(department_scope, int) else []
    conditions.extend(
        db.and_(Student.department_id.is_(None), db.func.lower(Student.department).in_(values))
    )
    return Student.query.filter(db.or_(*conditions))


def _visible_students():
    if has_role(current_user, Role.SUPER_ADMIN):
        return Student.query.order_by(Student.id).all()
    if has_role(current_user, Role.FACULTY):
        profile = Faculty.query.filter_by(user_id=current_user.id).first()
        if not profile:
            return None
        section_ids = {session.section_id for session in _faculty_class_sessions(profile)}
        if not section_ids:
            return []
        students = Student.query.filter(Student.section_id.in_(section_ids)).order_by(Student.id).all()
        return [student for student in students if _faculty_can_read_student(student)]
    if has_role(current_user, Role.HOD, Role.DEPARTMENT_ADMIN):
        scope = department_id_for(current_user)
        if not scope or not can_access(current_user, 'student:read', department_id=scope):
            return None
        return _students_for_department(scope).all()
    return None


def _subject_department_scope(subject):
    value = str(subject.department or '').strip()
    if not value:
        return None
    matches = Department.query.filter(db.or_(
        db.func.lower(db.func.trim(Department.code)) == value.casefold(),
        db.func.lower(db.func.trim(Department.name)) == value.casefold(),
    )).all()
    if len(matches) > 1:
        return None
    return matches[0].id if matches else value


def _faculty_session_is_authorized(profile, class_session, permission='student:read', *, subject=None, date_value=None):
    if not class_session or class_session.faculty_id != profile.id:
        return False
    assigned_subject = subject or class_session.subject
    if (
        not assigned_subject or assigned_subject.id != class_session.subject_id
        or assigned_subject.faculty_id != profile.id
    ):
        return False
    section = class_session.section
    department_id = section.batch.course.department_id if section and section.batch and section.batch.course else None
    section_department_values = _scope_department_values(department_id)
    subject_department_values = _scope_department_values(_subject_department_scope(assigned_subject))
    if (
        not department_id or not section_department_values.intersection(subject_department_values)
        or not can_access(current_user, permission, department_id=department_id)
    ):
        return False
    return date_value is None or class_session.weekday == date_value.weekday()


def _faculty_class_sessions(profile, permission='student:read', *, subject_id=None):
    query = ClassSession.query.join(Subject, ClassSession.subject_id == Subject.id).filter(
        ClassSession.faculty_id == profile.id, Subject.faculty_id == profile.id,
    )
    if subject_id is not None:
        query = query.filter(ClassSession.subject_id == subject_id)
    return [
        item for item in query.order_by(ClassSession.weekday, ClassSession.starts_at).all()
        if _faculty_session_is_authorized(profile, item, permission)
    ]


def _faculty_can_read_student(student, permission='student:read', *, class_session=None, subject=None, date_value=None):
    if not has_role(current_user, Role.FACULTY) or student.section_id is None:
        return False
    profile = Faculty.query.filter_by(user_id=current_user.id).first()
    if not profile:
        return False
    section = db.session.get(Section, student.section_id)
    section_department_id = (
        section.batch.course.department_id if section and section.batch and section.batch.course else None
    )
    if not section_department_id or not _scope_department_values(
        _student_department_scope(student)
    ).intersection(_scope_department_values(section_department_id)):
        return False
    sessions = [class_session] if class_session is not None else ClassSession.query.filter_by(
        section_id=student.section_id, faculty_id=profile.id,
    ).all()
    return any(
        item.section_id == student.section_id
        and _faculty_session_is_authorized(
            profile, item, permission, subject=subject, date_value=date_value,
        )
        for item in sessions
    )


def _can_read_student(student):
    if has_role(current_user, Role.SUPER_ADMIN):
        return True
    if has_role(current_user, Role.STUDENT):
        return can_access(
            current_user, 'attendance:read_own', resource_owner_id=student.user_id,
        )
    if has_role(current_user, Role.FACULTY):
        return _faculty_can_read_student(student)
    scope = _student_department_scope(student)
    return scope is not None and can_access(current_user, 'student:read', department_id=scope)


def _can_access_attendance(record, permission='attendance:read'):
    if has_role(current_user, Role.SUPER_ADMIN):
        return True
    student = record.student
    if has_role(current_user, Role.STUDENT):
        return permission == 'attendance:read' and can_access(
            current_user, 'attendance:read_own', resource_owner_id=student.user_id,
        )
    if has_role(current_user, Role.FACULTY):
        faculty = Faculty.query.filter_by(user_id=current_user.id).first()
        if not faculty or not record.session:
            return False
        faculty_permission = 'attendance:read' if permission == 'attendance:read' else 'attendance:mark'
        return _faculty_can_read_student(
            student, faculty_permission, class_session=record.session,
            subject=record.subject, date_value=record.date,
        )
    scope = _student_department_scope(student)
    return scope is not None and can_access(current_user, permission, department_id=scope)


@api_bp.route('/students')
@login_required
def students():
    students = _visible_students()
    if students is None:
        return jsonify({'error': 'Forbidden'}), 403
    data = []
    for student in students:
        records = Attendance.query.filter_by(student_id=student.id).all()
        if has_role(current_user, Role.FACULTY):
            records = [record for record in records if _can_access_attendance(record)]
        present = sum(1 for item in records if item.status == 'PRESENT')
        total = len(records)
        data.append({
            'id': student.id,
            'name': student.name,
            'register_number': student.register_number,
            'department': student.department,
            'attendance': calculate_percentage(present, total),
        })
    return jsonify(data)


@api_bp.route('/dashboard/stats')
@login_required
def dashboard_stats():
    if has_role(current_user, Role.MANAGEMENT):
        return jsonify({
            'departments': Department.query.count(),
            'total_students': Student.query.count(),
            'classes_conducted': Attendance.query.count(),
        })
    students = _visible_students()
    if students is None:
        return jsonify({'error': 'Forbidden'}), 403
    total_students = len(students)
    student_ids = [student.id for student in students]
    total_records = 0
    average = 0
    students_above_75 = 0
    if students:
        values = []
        for student in students:
            records = Attendance.query.filter_by(student_id=student.id).all()
            if has_role(current_user, Role.FACULTY):
                records = [record for record in records if _can_access_attendance(record)]
            total_records += len(records)
            present = sum(1 for item in records if item.status == 'PRESENT')
            total = len(records)
            percentage = calculate_percentage(present, total)
            values.append(percentage)
            students_above_75 += percentage >= 75
        average = round(sum(values) / len(values), 2)
    if not has_role(current_user, Role.FACULTY):
        total_records = Attendance.query.filter(Attendance.student_id.in_(student_ids)).count() if student_ids else 0
    return jsonify({
        'total_students': total_students,
        'average_attendance': average,
        'classes_conducted': total_records,
        'students_above_75': students_above_75,
    })


@api_bp.route('/prediction/student/<int:student_id>')
@login_required
def prediction_student(student_id):
    student = db.session.get(Student, student_id)
    if not student:
        return jsonify({'error': 'Student not found'}), 404
    if not _can_read_student(student):
        return jsonify({'error': 'Forbidden'}), 403
    records = Attendance.query.filter_by(student_id=student.id).order_by(Attendance.date, Attendance.id).all()
    if has_role(current_user, Role.FACULTY):
        records = [record for record in records if _can_access_attendance(record)]
    result, _history = record_student_prediction(student, records)
    return jsonify({'student': student.name, **result})


@api_bp.route('/prediction/student/<int:student_id>/history')
@login_required
def prediction_student_history(student_id):
    student = db.session.get(Student, student_id)
    if not student:
        return jsonify({'error': 'Student not found'}), 404
    if not _can_read_student(student):
        return jsonify({'error': 'Forbidden'}), 403
    history = PredictionHistory.query.filter_by(student_id=student.id).order_by(
        PredictionHistory.created_at.desc(), PredictionHistory.id.desc(),
    ).limit(50).all()
    return jsonify([{
        'created_at': item.created_at.isoformat() if item.created_at else None,
        'current_percentage': item.current_percentage,
        'predicted_percentage': item.predicted_percentage,
        'target_percentage': DEFAULT_TARGET_PERCENTAGE,
        'required_continuous_classes': required_classes_to_attend(
            item.current_percentage, DEFAULT_TARGET_PERCENTAGE,
            (item.feature_snapshot or {}).get('attendance_records', 0),
        ),
        'risk': risk_level(item.predicted_percentage),
    } for item in history])


@api_bp.route('/attendance/student/<int:student_id>')
@login_required
def attendance_student(student_id):
    student = db.session.get(Student, student_id)
    if not student:
        return jsonify({'error': 'Student not found'}), 404
    if not _can_read_student(student):
        return jsonify({'error': 'Forbidden'}), 403
    records = Attendance.query.filter_by(student_id=student_id).all()
    if has_role(current_user, Role.FACULTY):
        records = [record for record in records if _can_access_attendance(record)]
    data = []
    for record in records:
        data.append({
            'id': record.id,
            'date': record.date.isoformat(),
            'subject': record.subject.subject_name,
            'status': record.status,
        })
    return jsonify(data)


@api_bp.route('/attendance/subject/<int:subject_id>')
@login_required
def attendance_subject(subject_id):
    subject = db.session.get(Subject, subject_id)
    if not subject:
        return jsonify({'error': 'Subject not found'}), 404
    scope = _subject_department_scope(subject)
    allowed = has_role(current_user, Role.SUPER_ADMIN) or can_access(
        current_user, 'attendance:read', department_id=scope
    )
    if has_role(current_user, Role.FACULTY):
        faculty = Faculty.query.filter_by(user_id=current_user.id).first()
        allowed = bool(faculty and subject.faculty_id == faculty.id and can_access(
            current_user, 'attendance:read', department_id=scope
        ))
    if not allowed:
        return jsonify({'error': 'Forbidden'}), 403
    records = Attendance.query.filter_by(subject_id=subject_id).all()
    records = [record for record in records if _can_access_attendance(record)]
    return jsonify({'subject': subject.subject_name, 'attendance': [{
        'student': item.student.name,
        'date': item.date.isoformat(),
        'status': item.status,
    } for item in records]})


@api_bp.route('/attendance', methods=['POST'])
@login_required
def create_attendance():
    payload = request.get_json() or {}
    student_id = payload.get('student_id')
    subject_id = payload.get('subject_id')
    date_value = payload.get('date')
    status = str(payload.get('status') or '').upper()
    if not all([student_id, subject_id, date_value, status]):
        return jsonify({'error': 'Missing required fields'}), 400
    if status not in {'PRESENT', 'ABSENT', 'LATE', 'EXCUSED', 'OD', 'LEAVE'}:
        return jsonify({'error': 'Invalid attendance status'}), 400
    try:
        parsed_date = datetime.strptime(date_value, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return jsonify({'error': 'Date must use YYYY-MM-DD format'}), 400
    student = db.session.get(Student, student_id)
    subject = db.session.get(Subject, subject_id)
    if not student or not subject:
        return jsonify({'error': 'Student or subject not found'}), 404
    session_id = None
    if payload.get('session_id') not in (None, ''):
        try:
            session_id = int(payload['session_id'])
        except (TypeError, ValueError):
            return jsonify({'error': 'Invalid class session'}), 400
        class_session = db.session.get(ClassSession, session_id)
        if (
            not class_session or class_session.subject_id != subject.id
            or class_session.weekday != parsed_date.weekday()
        ):
            return jsonify({'error': 'Class session must match the subject and scheduled weekday'}), 400
        if student.section_id != class_session.section_id:
            return jsonify({'error': 'Student is not enrolled in the selected class session section'}), 400
    elif has_role(current_user, Role.FACULTY):
        return jsonify({'error': 'Faculty attendance requires an assigned class session'}), 403
    student_scope = _student_department_scope(student)
    subject_scope = _subject_department_scope(subject)
    if not _scope_department_values(student_scope).intersection(_scope_department_values(subject_scope)):
        return jsonify({'error': 'Student and subject must belong to the same department'}), 400
    if has_role(current_user, Role.FACULTY):
        faculty = Faculty.query.filter_by(user_id=current_user.id).first()
        if (
            not faculty or subject.faculty_id != faculty.id or not session_id
            or not _faculty_can_read_student(
                student, 'attendance:mark', class_session=class_session,
                subject=subject, date_value=parsed_date,
            )
        ):
            return jsonify({'error': 'Forbidden'}), 403
    elif not has_role(current_user, Role.SUPER_ADMIN) and not can_access(
        current_user, 'attendance:manage', department_id=student_scope
    ):
        return jsonify({'error': 'Forbidden'}), 403
    existing = Attendance.query.filter_by(
        student_id=student_id, subject_id=subject_id, date=parsed_date, session_id=session_id,
    ).first()
    if existing:
        return jsonify({'error': 'Duplicate attendance record already exists'}), 409
    record = Attendance(
        student_id=student_id, subject_id=subject_id, date=parsed_date, status=status,
        session_id=session_id, marked_by=current_user.id,
    )
    db.session.add(record)
    db.session.commit()
    return jsonify({'message': 'Attendance created', 'id': record.id}), 201


@api_bp.route('/attendance/<int:attendance_id>', methods=['PUT'])
@login_required
def update_attendance(attendance_id):
    record = db.session.get(Attendance, attendance_id)
    if not record:
        return jsonify({'error': 'Attendance record not found'}), 404
    if not _can_access_attendance(record, 'attendance:manage'):
        return jsonify({'error': 'Forbidden'}), 403
    payload = request.get_json() or {}
    if 'status' in payload:
        status = str(payload['status']).upper()
        if status not in {'PRESENT', 'ABSENT', 'LATE', 'EXCUSED', 'OD', 'LEAVE'}:
            return jsonify({'error': 'Invalid attendance status'}), 400
        if has_role(current_user, Role.FACULTY):
            created_at = record.created_at
            window_hours = max(0, current_app.config.get('ATTENDANCE_CORRECTION_WINDOW_HOURS', 24))
            is_expired = created_at is None or datetime.utcnow() - created_at > timedelta(hours=window_hours)
            if is_expired and status != record.status:
                reason = str(payload.get('reason') or '').strip()
                if not 10 <= len(reason) <= 1000:
                    return jsonify({'error': 'A correction reason between 10 and 1000 characters is required'}), 400
                if AttendanceCorrectionRequest.query.filter_by(
                    attendance_id=record.id, status='PENDING',
                ).first():
                    return jsonify({'error': 'A correction request is already pending'}), 409
                correction = AttendanceCorrectionRequest(
                    attendance_id=record.id,
                    previous_status=record.status,
                    new_status=status,
                    reason=reason,
                    requested_by=current_user.id,
                    requester_role=role_for(current_user).value,
                )
                db.session.add(correction)
                db.session.commit()
                return jsonify({'message': 'Correction submitted for review', 'request_id': correction.id}), 202
        record.status = status
    db.session.commit()
    return jsonify({'message': 'Attendance updated'})


@api_bp.route('/attendance/<int:attendance_id>', methods=['DELETE'])
@login_required
def delete_attendance(attendance_id):
    record = db.session.get(Attendance, attendance_id)
    if not record:
        return jsonify({'error': 'Attendance record not found'}), 404
    if not _can_access_attendance(record, 'attendance:manage'):
        return jsonify({'error': 'Forbidden'}), 403
    db.session.delete(record)
    db.session.commit()
    return jsonify({'message': 'Attendance deleted'})


@api_bp.route('/import-attendance', methods=['POST'])
@login_required
def import_attendance():
    if not has_role(current_user, Role.SUPER_ADMIN):
        return jsonify({'error': 'Forbidden'}), 403
    try:
        file = request.files.get('file')
        if file is None:
            raise ImportValidationError('No uploaded CSV file.')
        payload = validate_csv_upload(file.stream, file.filename or '')
    except ImportValidationError as error:
        db.session.rollback()
        return jsonify({'error': str(error)}), 400
    stage = create_import_stage(current_user.id, payload)
    return jsonify({
        'stage_id': stage.id, 'status': stage.status, 'expires_at': stage.expires_at.isoformat(),
        'preview_url': url_for('api.import_attendance_preview', stage_id=stage.id),
        'commit_url': url_for('api.commit_import', stage_id=stage.id),
        **import_preview_data(stage),
    }), 202


@api_bp.route('/import-attendance/<int:stage_id>')
@login_required
def import_attendance_preview(stage_id):
    if not has_role(current_user, Role.SUPER_ADMIN):
        return jsonify({'error': 'Forbidden'}), 403
    stage = AttendanceImportStage.query.filter_by(id=stage_id, owner_id=current_user.id).first()
    if not stage:
        return jsonify({'error': 'Import preview not found'}), 404
    if stage.expires_at <= datetime.utcnow() and stage.status == 'preview':
        expire_stage(stage)
    if stage.status != 'preview' or not stage.payload:
        return jsonify({'error': 'Import preview is no longer available.', 'status': stage.status}), 410
    return jsonify({
        'stage_id': stage.id, 'status': stage.status, 'expires_at': stage.expires_at.isoformat(),
        **import_preview_data(stage),
    })


@api_bp.route('/import-attendance/<int:stage_id>/commit', methods=['POST'])
@login_required
def commit_import(stage_id):
    if not has_role(current_user, Role.SUPER_ADMIN):
        return jsonify({'error': 'Forbidden'}), 403
    stage = AttendanceImportStage.query.filter_by(id=stage_id, owner_id=current_user.id).first()
    if not stage:
        return jsonify({'error': 'Import preview not found'}), 404
    payload = request.get_json(silent=True) or request.form
    expected_phrase = f'COMMIT ATTENDANCE IMPORT {stage.id}'
    if str(payload.get('confirmation_phrase', '')).strip() != expected_phrase:
        return jsonify({'error': f'Type the confirmation phrase: {expected_phrase}'}), 400
    result, imported = commit_import_stage(stage, current_user.id)
    if result == 'committed':
        return jsonify({'status': result, 'imported': imported}), 201
    status = 410 if result == 'expired' else 409
    return jsonify({'error': 'Import preview expired.' if result == 'expired' else 'Import preview is unavailable or attendance changed after preview.', 'status': result}), status


@api_bp.route('/import-attendance/<int:stage_id>', methods=['DELETE'])
@login_required
def cancel_import(stage_id):
    if not has_role(current_user, Role.SUPER_ADMIN):
        return jsonify({'error': 'Forbidden'}), 403
    stage = AttendanceImportStage.query.filter_by(id=stage_id, owner_id=current_user.id).first()
    if not stage:
        return jsonify({'error': 'Import preview not found'}), 404
    if stage.status == 'preview':
        stage.status = 'cancelled'
        stage.payload = None
        db.session.commit()
    return jsonify({'status': stage.status}), 200


@api_bp.route('/export/attendance.csv')
@login_required
def export_attendance():
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(['register_number', 'student_name', 'department', 'subject_code', 'subject_name', 'date', 'status'])
    query = Attendance.query.join(Student).join(Subject)
    if has_role(current_user, Role.STUDENT):
        student = Student.query.filter_by(user_id=current_user.id).first()
        query = query.filter(Attendance.student_id == student.id) if student else query.filter(Attendance.id == -1)
    elif has_role(current_user, Role.FACULTY):
        faculty = Faculty.query.filter_by(user_id=current_user.id).first()
        if not faculty or not record.session:
            return False
        faculty_permission = 'attendance:read' if permission == 'attendance:read' else 'attendance:mark'
        return _faculty_can_read_student(
            student, faculty_permission, class_session=record.session,
            subject=record.subject, date_value=record.date,
        )
        if visible_students is None or not can_access(
            current_user, 'report:read', department_id=department_id_for(current_user)
        ):
            return jsonify({'error': 'Forbidden'}), 403
        query = query.filter(Attendance.student_id.in_([item.id for item in visible_students]))
    elif not has_role(current_user, Role.SUPER_ADMIN):
        return jsonify({'error': 'Forbidden'}), 403
    records = query.order_by(Attendance.date, Attendance.id).all()
    if has_role(current_user, Role.FACULTY):
        records = [record for record in records if _can_access_attendance(record)]
    for record in records:
        writer.writerow([
            record.student.register_number, record.student.name, record.student.department,
            record.subject.subject_code, record.subject.subject_name, record.date.isoformat(), record.status,
        ])
    return Response(output.getvalue(), mimetype='text/csv', headers={
        'Content-Disposition': 'attachment; filename=attendance-report.csv'
    })


@api_bp.route('/export/attendance.<string:file_format>')
@login_required
def export_attendance_file(file_format):
    if file_format not in {'xlsx', 'pdf'}:
        return jsonify({'error': 'Supported attendance report formats are XLSX and PDF.'}), 404
    if not has_role(
        current_user, Role.SUPER_ADMIN, Role.MANAGEMENT, Role.HOD,
        Role.DEPARTMENT_ADMIN, Role.FACULTY, Role.STUDENT,
    ):
        return jsonify({'error': 'Forbidden'}), 403

    if has_role(current_user, Role.MANAGEMENT):
        if not can_access(current_user, 'report:read'):
            return jsonify({'error': 'Forbidden'}), 403
        rows = _management_attendance_report()
        headers = ['Department', 'Students with records', 'Attendance records', 'Present records', 'Attendance %']
        title = 'College attendance aggregate'
    else:
        student_ids = None
        faculty_id = None
        if has_role(current_user, Role.STUDENT):
            student = Student.query.filter_by(user_id=current_user.id).first()
            if not student or not can_access(
                current_user, 'attendance:read_own', resource_owner_id=student.user_id,
            ):
                return jsonify({'error': 'Forbidden'}), 403
            student_ids = [student.id]
            title = 'Student attendance report'
        elif has_role(current_user, Role.FACULTY):
            faculty = Faculty.query.filter_by(user_id=current_user.id).first()
            department_scope = department_id_for(current_user) or (faculty.department if faculty else None)
            if not faculty or not department_scope or not can_access(
                current_user, 'report:read', department_id=department_scope,
            ):
                return jsonify({'error': 'Forbidden'}), 403
            visible_students = _visible_students()
            if visible_students is None:
                return jsonify({'error': 'Forbidden'}), 403
            student_ids = [item.id for item in visible_students]
            faculty_id = faculty.id
            title = 'Faculty assigned attendance report'
        elif has_role(current_user, Role.HOD, Role.DEPARTMENT_ADMIN):
            department_scope = department_id_for(current_user)
            if not department_scope or not can_access(
                current_user, 'report:read', department_id=department_scope,
            ):
                return jsonify({'error': 'Forbidden'}), 403
            visible_students = _visible_students()
            if visible_students is None:
                return jsonify({'error': 'Forbidden'}), 403
            student_ids = [item.id for item in visible_students]
            title = 'Department attendance report'
        elif has_role(current_user, Role.SUPER_ADMIN):
            title = 'Institution attendance report'
        else:
            return jsonify({'error': 'Forbidden'}), 403

        query = Attendance.query.join(Student).join(Subject)
        if student_ids is not None:
            query = query.filter(Attendance.student_id.in_(student_ids)) if student_ids else query.filter(Attendance.id == -1)
        if faculty_id is not None:
            query = query.filter(Subject.faculty_id == faculty_id)
        records = query.order_by(Attendance.date, Attendance.id).limit(5001).all()
        if has_role(current_user, Role.FACULTY):
            records = [record for record in records if _can_access_attendance(record)]
        if len(records) > 5000:
            return jsonify({'error': 'Report exceeds the 5,000-row export limit; narrow the report scope.'}), 413
        headers = ['Register number', 'Student', 'Department', 'Subject code', 'Subject', 'Date', 'Status']
        rows = [[
            record.student.register_number, record.student.name, record.student.department,
            record.subject.subject_code, record.subject.subject_name, record.date.isoformat(), record.status,
        ] for record in records]

    if file_format == 'xlsx':
        content = _attendance_xlsx(title, headers, rows)
        return Response(content, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={
            'Content-Disposition': 'attachment; filename=attendance-report.xlsx',
        })
    content = _attendance_pdf(title, headers, rows)
    return Response(content, mimetype='application/pdf', headers={
        'Content-Disposition': 'attachment; filename=attendance-report.pdf',
    })


def _management_attendance_report():
    present_count = db.func.sum(db.case((Attendance.status == 'PRESENT', 1), else_=0))
    department_label = db.func.coalesce(Department.code, Student.department, 'Unassigned')
    query = db.session.query(
        department_label.label('department'),
        db.func.count(db.distinct(Student.id)).label('students'),
        db.func.count(Attendance.id).label('records'),
        present_count.label('present'),
    ).select_from(Attendance).join(Student).outerjoin(Department, Student.department_id == Department.id)
    aggregates = query.group_by(department_label).order_by(department_label).limit(5000).all()
    return [[
        item.department, item.students, item.records, item.present or 0,
        round((item.present or 0) * 100 / item.records, 2) if item.records else 0.0,
    ] for item in aggregates]


def _attendance_xlsx(title, headers, rows):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Attendance'
    sheet.append([title])
    sheet.append(headers)
    for row in rows:
        sheet.append([
            "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')) else value
            for value in row
        ])
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A2:{sheet.cell(row=2, column=len(headers)).column_letter}{max(2, len(rows) + 2)}'
    for cell in sheet[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    for column in sheet.columns:
        width = min(36, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _attendance_pdf(title, headers, rows):
    from html import escape
    from io import BytesIO

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(letter), rightMargin=0.35 * inch,
        leftMargin=0.35 * inch, topMargin=0.45 * inch, bottomMargin=0.45 * inch,
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('ReportCell', parent=styles['BodyText'], fontSize=7, leading=9, alignment=TA_LEFT)
    data = [[Paragraph(escape(str(value)), cell_style) for value in headers]]
    data.extend([[Paragraph(escape(str(value)), cell_style) for value in row] for row in rows])
    table = Table(data, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    document.build([Paragraph(escape(title), styles['Title']), Spacer(1, 10), table])
    return output.getvalue()
